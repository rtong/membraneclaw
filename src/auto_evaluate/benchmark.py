from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .io_utils import read_json, sha256_file, utc_now, write_json


REQUIRED_SHEETS = ("题目_Q", "分步答案_A", "评价标准")


@dataclass(frozen=True)
class BenchmarkSource:
    case_id: str
    path: Path
    task_family: str


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip()


def sheet_to_text(sheet) -> str:
    """Render populated worksheet rows as stable, model-readable text."""
    lines: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [_display(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if not any(cells):
            if lines and lines[-1] != "":
                lines.append("")
            continue
        lines.append(" | ".join(cell for cell in cells if cell))
    while lines and not lines[-1]:
        lines.pop()
    return "\n".join(lines)


def _first_nonempty(sheet) -> str:
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            text = _display(value)
            if text:
                return text
    return sheet.title


def _parse_rubric(sheet) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    header_row = None
    for idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = [_display(value) for value in row]
        if values and values[0] in {"步骤", "能力层"} and "分值" in values:
            header_row = idx
            break
    if header_row is None:
        raise ValueError(f"{sheet.title}: cannot locate rubric header row")

    steps: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_id = _display(row[0] if row else None)
        score = row[6] if len(row) > 6 else None
        if not raw_id or not isinstance(score, (int, float)):
            if steps:
                break
            continue
        steps.append(
            {
                "step_id": len(steps) + 1,
                "step_label": raw_id,
                "reasoning_or_calculation": _display(row[1] if len(row) > 1 else None),
                "inputs": _display(row[2] if len(row) > 2 else None),
                "key_outputs": _display(row[3] if len(row) > 3 else None),
                "full_credit": _display(row[4] if len(row) > 4 else None),
                "common_failures": _display(row[5] if len(row) > 5 else None),
                "max_score": float(score),
            }
        )

    global_criteria: list[dict[str, str]] = []
    in_supplement = False
    for row in sheet.iter_rows(values_only=True):
        values = [_display(value) for value in row]
        if values and values[0] == "总体判分补充":
            in_supplement = True
            continue
        if not in_supplement or not values or not values[0]:
            continue
        if len(values) > 1 and values[1]:
            global_criteria.append(
                {
                    "name": values[0],
                    "criterion": values[1],
                    "penalty_guidance": values[2] if len(values) > 2 else "",
                }
            )
    return steps, global_criteria


def import_workbook(source: BenchmarkSource) -> dict[str, Any]:
    if not source.path.exists():
        raise FileNotFoundError(f"Benchmark source not found: {source.path}")
    workbook = load_workbook(source.path, data_only=False, read_only=True)
    missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"{source.path.name}: missing sheets {missing}")

    question = workbook["题目_Q"]
    answer = workbook["分步答案_A"]
    rubric_sheet = workbook["评价标准"]
    steps, global_criteria = _parse_rubric(rubric_sheet)

    result = {
        "schema_version": "1.0",
        "case_id": source.case_id,
        "task_family": source.task_family,
        "title": _first_nonempty(question),
        "question_prompt": sheet_to_text(question),
        "reference_answer": sheet_to_text(answer),
        "rubric": {
            "total_points": sum(step["max_score"] for step in steps),
            "steps": steps,
            "global_criteria": global_criteria,
        },
        "source": {
            "filename": source.path.name,
            "absolute_path": str(source.path.resolve()),
            "sha256": sha256_file(source.path),
            "sheets": list(REQUIRED_SHEETS),
            "imported_at": utc_now(),
        },
    }
    validate_benchmark(result)
    return result


def validate_benchmark(benchmark: dict[str, Any]) -> None:
    required = ("case_id", "task_family", "title", "question_prompt", "reference_answer", "rubric")
    missing = [key for key in required if not benchmark.get(key)]
    if missing:
        raise ValueError(f"Benchmark missing required fields: {missing}")
    steps = benchmark["rubric"].get("steps") or []
    if not steps:
        raise ValueError(f"{benchmark['case_id']}: rubric has no steps")
    ids = [step["step_id"] for step in steps]
    if ids != list(range(1, len(ids) + 1)):
        raise ValueError(f"{benchmark['case_id']}: rubric steps are not consecutive: {ids}")
    total = sum(float(step["max_score"]) for step in steps)
    declared = float(benchmark["rubric"].get("total_points", 0))
    if abs(total - declared) > 1e-9:
        raise ValueError(f"{benchmark['case_id']}: rubric total mismatch {total} != {declared}")
    if abs(total - 100.0) > 1e-9:
        raise ValueError(f"{benchmark['case_id']}: expected a 100-point rubric, got {total}")


def load_sources(config_path: Path, project_root: Path) -> list[BenchmarkSource]:
    config = read_json(config_path)
    sources = []
    seen_case_ids: set[str] = set()
    for item in config.get("sources", []):
        case_id = item["case_id"]
        if case_id in seen_case_ids:
            raise ValueError(f"Duplicate benchmark case_id in {config_path}: {case_id}")
        seen_case_ids.add(case_id)
        raw_path = Path(item["path"])
        path = raw_path if raw_path.is_absolute() else (project_root / raw_path)
        sources.append(
            BenchmarkSource(
                case_id=case_id,
                path=path.resolve(),
                task_family=item["task_family"],
            )
        )
    if not sources:
        raise ValueError(f"No benchmark sources configured in {config_path}")
    return sources


def import_all(config_path: Path, output_dir: Path, project_root: Path) -> list[dict[str, Any]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    benchmarks = [import_workbook(source) for source in load_sources(config_path, project_root)]
    for benchmark in benchmarks:
        write_json(output_dir / f"{benchmark['case_id']}.json", benchmark)
    write_json(
        output_dir / "index.json",
        {
            "schema_version": "1.0",
            "generated_at": utc_now(),
            "cases": [
                {
                    "case_id": item["case_id"],
                    "task_family": item["task_family"],
                    "title": item["title"],
                    "file": f"{item['case_id']}.json",
                    "source_sha256": item["source"]["sha256"],
                }
                for item in benchmarks
            ],
        },
    )
    return benchmarks


def iter_benchmarks(directory: Path) -> Iterable[dict[str, Any]]:
    index = read_json(directory / "index.json")
    for case in index.get("cases", []):
        benchmark = read_json(directory / case["file"])
        validate_benchmark(benchmark)
        yield benchmark
